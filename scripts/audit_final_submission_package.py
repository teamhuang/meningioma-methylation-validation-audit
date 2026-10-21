from pathlib import Path
import re,gzip,csv
from docx import Document
from pypdf import PdfReader

pkg=Path(r'E:\脑膜瘤npj\npjDM_submission_candidate_v6')
md=(pkg/'manuscript_npjDM_submission_candidate_v6.md').read_text(encoding='utf-8-sig')
abs_txt=md.split('## Abstract',1)[1].split('---',1)[0]
words=re.findall(r"\b[\w]+(?:[’'-][\w]+)*\b",abs_txt,flags=re.UNICODE)
refs_block=md.split('## References',1)[1].split('## Figure legends',1)[0]
refs=[int(m.group(1)) for line in refs_block.splitlines() if (m:=re.match(r'\s*(\d+)\.',line))]
figs=[int(x) for x in re.findall(r'\*\*Figure (\d+)\.',md)]
sfigs=[int(x) for x in re.findall(r'\*\*Supplementary Figure S(\d+)\.',md)]
d=Document(pkg/'manuscript_npjDM_submission_candidate_v6.docx')
si=PdfReader(str(pkg/'Supplementary_Information.pdf'))
print('Abstract words:',len(words))
print('Abstract whitespace tokens:',len(abs_txt.split()))
print('Reference numbering:',refs)
print('Main figure captions:',figs)
print('Supplementary figure captions:',sfigs)
print('DOCX tables/embedded figures:',len(d.tables),len(d.inline_shapes))
print('Supplementary PDF pages:',len(si.pages))
print('Author placeholders:',md.count('[AUTHOR'))
for name in ['Supplementary_Data_1_corrected_panel_results.csv.gz','Supplementary_Data_2_longitudinal_mapping.csv','Supplementary_Data_3_ucsf_identity_linkage.csv']:
    p=pkg/name
    if p.suffixes[-2:]==['.csv','.gz']:
        with gzip.open(p,'rt',encoding='utf-8-sig',newline='') as f: rows=sum(1 for _ in csv.reader(f))-1
    else:
        with p.open('r',encoding='utf-8-sig',newline='') as f: rows=sum(1 for _ in csv.reader(f))-1
    print(name, 'data rows:',rows)
print('No extraneous duplicate supplementary data names:',not any(p.name in {'Supplementary_Data_1_longitudinal_mapping.csv','Supplementary_Data_2_ucsf_identity_linkage.csv','Supplementary_Data_3_corrected_panel_results.csv.gz'} for p in pkg.iterdir()))
