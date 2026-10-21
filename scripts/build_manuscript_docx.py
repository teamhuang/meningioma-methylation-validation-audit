from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r'E:\脑膜瘤npj')
PKG = ROOT / 'npjDM_submission_candidate_v6'
MD = PKG / 'manuscript_npjDM_submission_candidate_v6.md'
OUT = PKG / 'manuscript_npjDM_submission_candidate_v6.docx'
FIGS = {i: PKG / f'Figure{i}.png' for i in range(1, 7)}

doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.0); section.bottom_margin = Cm(2.0)
section.left_margin = Cm(2.25); section.right_margin = Cm(2.25)

normal = doc.styles['Normal']
normal.font.name = 'Arial'; normal.font.size = Pt(10.5)
normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.12
for sn in ['Heading 1','Heading 2','Heading 3']:
    st = doc.styles[sn]; st.font.name = 'Arial'; st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    st.font.color.rgb = __import__('docx').shared.RGBColor(0,0,0)

def add_page_num(sec):
    p = sec.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(); fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); r._r.addnext(fld)
add_page_num(section)

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def set_cell_margins(cell, top=55, start=65, bottom=55, end=65):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); mar = OxmlElement('w:tcMar')
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=OxmlElement('w:'+m); node.set(qn('w:w'),str(v)); node.set(qn('w:type'),'dxa'); mar.append(node)
    tcPr.append(mar)

def add_inline(p, text):
    # Small markdown subset: bold and italic. All other content remains literal text.
    parts = re.split(r'(\*\*.*?\*\*|(?<!\*)\*[^*]+\*(?!\*))', text)
    for s in parts:
        if not s: continue
        if s.startswith('**') and s.endswith('**'):
            run=p.add_run(s[2:-2]); run.bold=True
        elif s.startswith('*') and s.endswith('*'):
            run=p.add_run(s[1:-1]); run.italic=True
        else:
            p.add_run(s)

def parse_table(lines, is_first=False):
    rows=[]
    for line in lines:
        cells=[c.strip() for c in line.strip().strip('|').split('|')]
        if cells and all(re.fullmatch(r':?-{3,}:?', c.replace(' ','')) for c in cells): continue
        rows.append(cells)
    if not rows: return
    ncol=max(map(len,rows)); rows=[r+['']*(ncol-len(r)) for r in rows]
    table=doc.add_table(rows=1, cols=ncol); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
    table.style='Table Grid'
    sec=doc.sections[-1]
    width=sec.page_width-sec.left_margin-sec.right_margin
    colw=width/ncol
    for row_i,vals in enumerate(rows):
        cells=table.rows[0].cells if row_i==0 else table.add_row().cells
        table.rows[row_i]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for j,val in enumerate(vals):
            cell=cells[j]; cell.width=colw; set_cell_margins(cell)
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
            add_inline(p,val)
            for run in p.runs: run.font.name='Arial'; run.font.size=Pt(7.1 if ncol>=6 else 8)
            if row_i==0:
                set_cell_shading(cell,'D9EAF7')
                for run in p.runs: run.bold=True
        if row_i==0:
            trPr=table.rows[0]._tr.get_or_add_trPr(); rep=OxmlElement('w:tblHeader'); rep.set(qn('w:val'),'true'); trPr.append(rep)
    doc.add_paragraph()

lines=MD.read_text(encoding='utf-8-sig').splitlines()
i=0; paragraph=[]; in_figleg=False; table_section=False
def flush():
    global paragraph
    if not paragraph: return
    text=' '.join(x.strip() for x in paragraph).strip()
    if not text: paragraph=[]; return
    p=doc.add_paragraph(); p.paragraph_format.keep_together=True
    if 'Jianhuang Huang' in text and 'Jianning Chen' in text:
        p.add_run(text.replace('\\*','*'))
    else:
        add_inline(p,text)
    for m in re.finditer(r'\*\*Figure\s+(\d+)\.', text):
        figno=int(m.group(1))
        if figno in FIGS and FIGS[figno].exists():
            p.paragraph_format.keep_with_next=True
            q=doc.add_paragraph(); q.alignment=WD_ALIGN_PARAGRAPH.CENTER; q.paragraph_format.space_after=Pt(2)
            q.add_run().add_picture(str(FIGS[figno]), width=Cm(16.2))
            q.paragraph_format.keep_together=True
    paragraph=[]

while i < len(lines):
    raw=lines[i]; s=raw.strip()
    if not s or s=='---': flush(); i+=1; continue
    if s.startswith('|'):
        flush(); tb=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            tb.append(lines[i]); i+=1
        parse_table(tb); continue
    if s.startswith('# '):
        flush(); p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after=Pt(14); add_inline(p,s[2:]);
        for r in p.runs: r.bold=True; r.font.size=Pt(16)
        i+=1; continue
    if s.startswith('## '):
        flush(); title=s[3:]
        if title=='Tables' and not table_section:
            sec=doc.add_section(WD_SECTION.NEW_PAGE); sec.orientation=1
            sec.page_width=Cm(29.7); sec.page_height=Cm(21.0)
            sec.top_margin=Cm(1.5); sec.bottom_margin=Cm(1.5); sec.left_margin=Cm(1.4); sec.right_margin=Cm(1.4)
            sec.footer.is_linked_to_previous=False
            add_page_num(sec); table_section=True
            p=doc.add_paragraph(); p.style='Heading 1'; p.add_run(title)
        else:
            p=doc.add_paragraph(); p.style='Heading 1'; add_inline(p,title)
        i+=1; continue
    if s.startswith('### '):
        flush(); p=doc.add_paragraph(); p.style='Heading 2'; add_inline(p,s[4:]); i+=1; continue
    # Figure legends become compact while the plate follows its own caption.
    paragraph.append(s)
    i+=1
flush()
doc.core_properties.title='Identity, training independence and selective prediction shape the transportability of a meningioma DNA methylation classifier'
doc.core_properties.author='Jianhuang Huang; Yangyang Lin; Qixiu Wang; Jianning Chen'
doc.save(OUT)
print(f'Created {OUT}')
